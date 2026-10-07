#!/usr/bin/env python3
"""Dump Office / PDF / text files to readable text so audit subagents never
have to open binaries.

Must run under the workspace venv, which carries the readers:
    <repo>/.venv/bin/python extract_files.py <dir> <out.md>

Writes one markdown file covering every file in <dir>. Unreadable files are
reported as such rather than skipped silently — an auditor needs to tell
"this is wrong" from "I could not read this".
"""
from __future__ import annotations

import sys
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import audit_common as ac  # noqa: E402

MAX_CHARS = 20000        # per file, to keep prompts bounded
MAX_ROWS = 200           # per worksheet


def _docx(p: Path) -> str:
    import docx
    d = docx.Document(str(p))
    out = [f"paragraphs: {len(d.paragraphs)}, tables: {len(d.tables)}"]
    for i, para in enumerate(d.paragraphs):
        if para.text.strip():
            style = para.style.name if para.style is not None else ""
            out.append(f"[p{i}]{f' ({style})' if style and style != 'Normal' else ''} "
                       f"{para.text.strip()}")
    for ti, t in enumerate(d.tables):
        out.append(f"\n[table {ti}] {len(t.rows)}x{len(t.columns)}")
        for r in t.rows[:MAX_ROWS]:
            out.append("  | " + " | ".join(c.text.strip() for c in r.cells))
    return "\n".join(out)


def _xlsx(p: Path) -> str:
    import openpyxl
    out = []
    for data_only in (True, False):
        wb = openpyxl.load_workbook(str(p), data_only=data_only)
        label = "VALUES" if data_only else "FORMULAS"
        out.append(f"=== {label} ===")
        for ws in wb.worksheets:
            out.append(f"\n[sheet: {ws.title}] dims={ws.dimensions} "
                       f"max_row={ws.max_row} max_col={ws.max_column}")
            for row in ws.iter_rows(max_row=min(ws.max_row or 1, MAX_ROWS)):
                cells = [("" if c.value is None else str(c.value)) for c in row]
                if any(cells):
                    out.append(f"  r{row[0].row}: " + " | ".join(cells))
        if data_only:
            out.append("")
    return "\n".join(out)


def _pptx(p: Path) -> str:
    import pptx
    pr = pptx.Presentation(str(p))
    out = [f"slides: {len(pr.slides)}, "
           f"size: {pr.slide_width}x{pr.slide_height} EMU"]
    for i, s in enumerate(pr.slides, 1):
        out.append(f"\n[slide {i}] layout={s.slide_layout.name}")
        for sh in s.shapes:
            kind = sh.shape_type
            pos = (f"@({sh.left},{sh.top}) {sh.width}x{sh.height}"
                   if sh.left is not None else "")
            txt = ""
            if sh.has_text_frame and sh.text_frame.text.strip():
                txt = " :: " + sh.text_frame.text.strip().replace("\n", " / ")
            out.append(f"  - {sh.shape_type and str(kind)} name={sh.name!r} "
                       f"{pos}{txt}")
            if getattr(sh, "has_table", False):
                for r in sh.table.rows:
                    out.append("      | " + " | ".join(
                        c.text.strip() for c in r.cells))
    return "\n".join(out)


def _pdf(p: Path) -> str:
    import pdfplumber
    out = []
    with pdfplumber.open(str(p)) as pdf:
        out.append(f"pages: {len(pdf.pages)}")
        for i, pg in enumerate(pdf.pages, 1):
            out.append(f"\n[page {i}] {pg.width:.0f}x{pg.height:.0f}pt")
            out.append((pg.extract_text() or "(no extractable text)").strip())
            for t in pg.extract_tables() or []:
                out.append("  [table]")
                for r in t[:MAX_ROWS]:
                    out.append("    | " + " | ".join(
                        "" if c is None else str(c).strip() for c in r))
    return "\n".join(out)


def _zip(p: Path) -> str:
    with zipfile.ZipFile(p) as z:
        names = z.namelist()
        out = [f"zip members: {len(names)}"]
        out += [f"  {n}" for n in names[:300]]
        return "\n".join(out)


def _text(p: Path) -> str:
    return p.read_text(encoding="utf-8", errors="replace")


HANDLERS = {
    ".docx": _docx, ".xlsx": _xlsx, ".xlsm": _xlsx, ".pptx": _pptx,
    ".pdf": _pdf, ".zip": _zip,
}
TEXTY = {".txt", ".md", ".csv", ".tsv", ".json", ".html", ".htm", ".css",
         ".js", ".py", ".r", ".sql", ".xml", ".yaml", ".yml", ".svg"}


def render(p: Path) -> str:
    ext = p.suffix.lower()
    try:
        if ext in HANDLERS:
            body = HANDLERS[ext](p)
        elif ext in TEXTY:
            body = _text(p)
        elif ext in {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp"}:
            try:
                from PIL import Image
                with Image.open(p) as im:
                    return (f"IMAGE {im.format} {im.size[0]}x{im.size[1]} "
                            f"mode={im.mode}\n(open the file directly to judge "
                            f"appearance — text extraction does not apply)")
            except Exception as e:                          # noqa: BLE001
                return f"IMAGE — could not read header: {type(e).__name__}: {e}"
        else:
            return f"UNSUPPORTED EXTENSION {ext} ({p.stat().st_size} bytes)"
    except Exception as e:                                  # noqa: BLE001
        return f"COULD NOT READ: {type(e).__name__}: {e}"
    if len(body) > MAX_CHARS:
        body = body[:MAX_CHARS] + f"\n… TRUNCATED at {MAX_CHARS} chars"
    return body


def main() -> int:
    if len(sys.argv) < 3:
        print(__doc__)
        return 2
    src, dest = Path(sys.argv[1]), Path(sys.argv[2])
    unit = dest.resolve().parent
    if unit.parent == ac.AUDIT.resolve():
        why = ac.locked(unit.name, "extract_files.py")
        if why:
            print(why)
            return 3
    if not src.exists():
        dest.write_text(f"# (no files staged at {src})\n", encoding="utf-8")
        return 0
    files = sorted(p for p in src.rglob("*") if p.is_file())
    out = [f"# Extracted content of {src.name}/ ({len(files)} files)", "",
           "Machine-extracted. Images and rendered appearance cannot be judged "
           "from this file — open the original for that.", ""]
    for p in files:
        out += [f"## {p.relative_to(src)}  ({p.stat().st_size} bytes)", "",
                "```", render(p), "```", ""]
    dest.parent.mkdir(parents=True, exist_ok=True)
    dest.write_text("\n".join(out), encoding="utf-8")
    print(f"  extracted {len(files)} file(s) -> {dest}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
